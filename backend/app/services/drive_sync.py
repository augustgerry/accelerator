"""
Google Drive integration. Files stay in Drive — only extracted text +
embeddings are stored locally. See README for OAuth setup.

First run opens a browser for the OAuth consent screen (credentials from
GOOGLE_OAUTH_CREDENTIALS_PATH); the resulting token is cached in token.json
next to it, so later runs don't need to log in again.
"""

import io
import os

from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload

from app.config import settings

SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]
TOKEN_PATH = os.path.join(
    os.path.dirname(settings.google_oauth_credentials_path) or ".", "token.json"
)

FOLDER_MIME = "application/vnd.google-apps.folder"
GOOGLE_DOC_MIME = "application/vnd.google-apps.document"
PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _get_credentials() -> Credentials:
    creds = None
    if os.path.exists(TOKEN_PATH):
        creds = Credentials.from_authorized_user_file(TOKEN_PATH, SCOPES)

    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            flow = InstalledAppFlow.from_client_secrets_file(
                settings.google_oauth_credentials_path, SCOPES
            )
            creds = flow.run_local_server(port=0)
        with open(TOKEN_PATH, "w") as f:
            f.write(creds.to_json())

    return creds


def _get_service():
    return build("drive", "v3", credentials=_get_credentials())


def _list_children(service, folder_id: str) -> list[dict]:
    children: list[dict] = []
    page_token = None
    query = f"'{folder_id}' in parents and trashed = false"
    while True:
        response = (
            service.files()
            .list(
                q=query,
                fields="nextPageToken, files(id, name, mimeType, modifiedTime)",
                pageToken=page_token,
                supportsAllDrives=True,
                includeItemsFromAllDrives=True,
                corpora="allDrives",
            )
            .execute()
        )
        children.extend(response.get("files", []))
        page_token = response.get("nextPageToken")
        if not page_token:
            break
    return children


def _list_files_recursive(service, folder_id: str, path_parts: list[str]) -> list[dict]:
    files: list[dict] = []
    for child in _list_children(service, folder_id):
        if child["mimeType"] == FOLDER_MIME:
            files.extend(_list_files_recursive(service, child["id"], path_parts + [child["name"]]))
        else:
            child["folderPath"] = "/".join(path_parts)
            files.append(child)
    return files


def list_drive_files(folder_id: str | None = None) -> list[dict]:
    """Recursively list files inside `folder_id or settings.google_drive_folder_id`,
    descending into every subfolder. Each file dict carries `folderPath` — the
    subfolder path (e.g. "Sulfindo/2024") it was found under, "" if directly
    in the root folder — for later mapping onto `Document.division`."""
    folder_id = folder_id or settings.google_drive_folder_id
    service = _get_service()
    return _list_files_recursive(service, folder_id, [])


def _download_media(service, file_id: str) -> bytes:
    request = service.files().get_media(fileId=file_id, supportsAllDrives=True)
    buffer = io.BytesIO()
    downloader = MediaIoBaseDownload(buffer, request)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    return buffer.getvalue()


def download_file_bytes(file_id: str) -> bytes:
    """Download the raw bytes of a Drive file (e.g. to re-serve a .docx template)."""
    service = _get_service()
    return _download_media(service, file_id)


def _format_docx_table_markdown(table) -> str:
    """Format a python-docx Table into clean Markdown table string."""
    rows = []
    for row in table.rows:
        row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        if any(row_cells):
            rows.append("| " + " | ".join(row_cells) + " |")
    if not rows:
        return ""
    header_cols = len(table.columns) if table.columns else max(1, rows[0].count("|") - 1)
    separator = "| " + " | ".join(["---"] * max(1, header_cols)) + " |"
    if len(rows) > 1:
        return "\n" + rows[0] + "\n" + separator + "\n" + "\n".join(rows[1:]) + "\n"
    return "\n" + rows[0] + "\n" + separator + "\n"


def _extract_docx_text_and_tables(file_bytes: bytes) -> str:
    """Extract paragraphs with heading levels AND tables in sequential document order."""
    from docx import Document as DocxDocument
    from docx.text.paragraph import Paragraph
    from docx.table import Table

    doc = DocxDocument(io.BytesIO(file_bytes))
    blocks = []
    for child in doc.element.body:
        if child.tag.endswith("p"):
            p = Paragraph(child, doc)
            text = p.text.strip()
            if text:
                style_name = (p.style.name or "").lower() if p.style else ""
                if "heading 1" in style_name:
                    blocks.append(f"# {text}")
                elif "heading 2" in style_name:
                    blocks.append(f"## {text}")
                elif "heading 3" in style_name:
                    blocks.append(f"### {text}")
                else:
                    blocks.append(text)
        elif child.tag.endswith("tbl"):
            tbl = Table(child, doc)
            tbl_md = _format_docx_table_markdown(tbl)
            if tbl_md:
                blocks.append(tbl_md)
    return "\n\n".join(blocks)


def fetch_and_extract_text(file_id: str) -> str:
    """Download file `file_id` and extract plain text based on its mimeType."""
    service = _get_service()
    meta = (
        service.files()
        .get(fileId=file_id, fields="mimeType, name", supportsAllDrives=True)
        .execute()
    )
    mime = meta["mimeType"]

    if mime == GOOGLE_DOC_MIME:
        from app.services.document_parser import clean_signature_blocks

        data = service.files().export(fileId=file_id, mimeType="text/plain").execute()
        clean_text, _ = clean_signature_blocks(data.decode("utf-8"))
        return clean_text

    if mime == PDF_MIME:
        from app.services.document_parser import extract_pdf_with_ocr_fallback

        text, _ = extract_pdf_with_ocr_fallback(_download_media(service, file_id))
        return text

    if mime == DOCX_MIME:
        from app.services.document_parser import clean_signature_blocks

        raw_docx = _extract_docx_text_and_tables(_download_media(service, file_id))
        clean_docx, _ = clean_signature_blocks(raw_docx)
        return clean_docx

    if mime == PPTX_MIME:
        from pptx import Presentation

        prs = Presentation(io.BytesIO(_download_media(service, file_id)))
        return "\n".join(
            shape.text_frame.text
            for slide in prs.slides
            for shape in slide.shapes
            if shape.has_text_frame
        )

    if mime.startswith("text/"):
        return _download_media(service, file_id).decode("utf-8")

    raise ValueError(f"Unsupported mimeType for file {file_id}: {mime}")
